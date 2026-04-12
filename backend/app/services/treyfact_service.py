"""
TreyFact-compatible Excel export.

TreyFact (Treyder Informática) accepts article imports via Excel.
Column names match the TreyFact import wizard field labels (Spanish).
"""
import io
import logging
from typing import List

logger = logging.getLogger(__name__)

# TreyFact import column layout
# Field name → (header label, column width)
TREYFACT_COLUMNS = [
    ("codigo",              "Código",                 16),
    ("nombre",              "Nombre",                 40),
    ("precio",              "Precio",                 12),
    ("iva",                 "% IVA",                  8),
    ("familia",             "Familia",                18),
    ("subfamilia",          "Subfamilia",             18),
    ("marca",               "Marca",                  16),
    ("ref_proveedor",       "Referencia Proveedor",   20),
    ("proveedor",           "Proveedor",              20),
    ("ean",                 "Código EAN",             16),
    ("margen",              "Margen Beneficio (%)",   18),
    ("pvp_con_iva",         "PVP con IVA",            12),
    ("coste_neto",          "Coste Neto",             12),
]


def generate_treyfact_excel(articles, supplier_name: str = "") -> bytes:
    """Generate a TreyFact-compatible Excel for article import.

    Args:
        articles: List of Article ORM objects
        supplier_name: Default supplier name (from document)

    Returns:
        Excel file bytes
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl not installed")
        return b""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Artículos TreyFact"

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2D6A9F", end_color="2D6A9F", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, (_, label, width) in enumerate(TREYFACT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TREYFACT_COLUMNS))}1"

    even_fill = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
    price_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    align_l = Alignment(horizontal="left",   vertical="center")
    align_r = Alignment(horizontal="right",  vertical="center")
    align_c = Alignment(horizontal="center", vertical="center")

    for row_idx, art in enumerate(articles, start=2):
        # Best article code: prefer codigo_principal, then fabricante, then proveedor
        codigo = (
            art.codigo_principal or
            art.codigo_fabricante or
            art.codigo_proveedor or
            ""
        )
        row = {
            "codigo":       codigo,
            "nombre":       art.descripcion or "",
            "precio":       round(art.pvp_sin_iva, 4),
            "iva":          int(art.iva_pct or 21),
            "familia":      art.familia or "",
            "subfamilia":   "",
            "marca":        art.codigo_fabricante or "",
            "ref_proveedor": art.codigo_proveedor or "",
            "proveedor":    supplier_name,
            "ean":          art.ean or "",
            "margen":       round(art.margen_pct, 2),
            "pvp_con_iva":  round(art.pvp_con_iva, 4),
            "coste_neto":   round(art.coste_neto_unitario, 4),
        }

        is_even = row_idx % 2 == 0

        for col_idx, (field, _, _) in enumerate(TREYFACT_COLUMNS, start=1):
            value = row[field]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin

            if field in ("precio", "pvp_con_iva", "coste_neto"):
                cell.fill = price_fill
                cell.number_format = '#,##0.0000 "€"'
                cell.alignment = align_r
            elif field == "margen":
                if is_even:
                    cell.fill = even_fill
                cell.number_format = '0.00"%"'
                cell.alignment = align_r
            elif field == "iva":
                if is_even:
                    cell.fill = even_fill
                cell.alignment = align_c
            elif field in ("nombre", "familia", "subfamilia", "proveedor"):
                if is_even:
                    cell.fill = even_fill
                cell.alignment = align_l
            else:
                if is_even:
                    cell.fill = even_fill
                cell.alignment = align_c

        ws.row_dimensions[row_idx].height = 16

    # Add a note sheet with instructions
    ws_notes = wb.create_sheet("Instrucciones")
    ws_notes["A1"] = "Instrucciones de importación en TreyFact"
    ws_notes["A1"].font = Font(bold=True, size=13)
    instructions = [
        "",
        "1. En TreyFact, ve a: Mantenimiento → Artículos → Importar artículos desde Excel",
        "2. Selecciona este fichero (.xlsx)",
        "3. Asegúrate de que la fila de inicio es la 2 (la 1 son los encabezados)",
        "4. Mapea cada columna de TreyFact con la columna correspondiente de este Excel:",
        "   - Código          → columna A",
        "   - Nombre          → columna B",
        "   - Precio (sin IVA)→ columna C",
        "   - % IVA           → columna D",
        "   - Familia         → columna E",
        "   - Marca           → columna G",
        "   - Ref. Proveedor  → columna H",
        "   - Proveedor       → columna I",
        "   - Código EAN      → columna J",
        "   - Margen (%)      → columna K",
        "5. Revisa y confirma la importación.",
        "",
        "Nota: el campo 'Precio' es PVP SIN IVA. TreyFact calcula el PVP con IVA automáticamente.",
    ]
    for i, line in enumerate(instructions, start=2):
        ws_notes[f"A{i}"] = line
    ws_notes.column_dimensions["A"].width = 70

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
