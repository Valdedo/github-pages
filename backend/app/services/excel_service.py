"""
Excel export service using openpyxl.
Generates a professional .xlsx file with all article data.
"""
import io
import json
import logging
from typing import List, Optional

logger = logging.getLogger(__name__)

# Column definitions: (field_name, header_label, width, number_format)
COLUMNS = [
    ("proveedor",               "Proveedor",              20, None),
    ("fecha_doc",               "Fecha Doc.",             14, None),
    ("num_doc",                 "Nº Albarán",             15, None),
    ("descripcion",             "Descripción",            40, None),
    ("cantidad",                "Cantidad",               10, "#,##0.###"),
    ("precio_unitario_bruto",   "P. Bruto Unit.",         14, "#,##0.00€"),
    ("descuento_1",             "Dto. 1 (%)",             10, "0.00%"),
    ("descuento_2",             "Dto. 2 (%)",             10, "0.00%"),
    ("descuento_3",             "Dto. 3 (%)",             10, "0.00%"),
    ("descuento_4",             "Dto. 4 (%)",             10, "0.00%"),
    ("coste_neto_unitario",     "Coste Neto Unit.",       14, "#,##0.0000€"),
    ("coste_neto_total",        "Coste Neto Total",       14, "#,##0.00€"),
    ("iva_pct",                 "IVA (%)",                8,  "0.00%"),
    ("recargo_pct",             "Rec. Equiv. (%)",        12, "0.00%"),
    ("margen_pct",              "Margen (%)",             10, "0.00%"),
    ("pvp_sin_iva",             "PVP sin IVA",            12, "#,##0.00€"),
    ("pvp_con_iva",             "PVP con IVA",            12, "#,##0.00€"),
    ("codigo_principal",        "Código Principal",       16, None),
    ("ean",                     "EAN/UPC",                16, None),
    ("codigo_fabricante",       "Cód. Fabricante",        16, None),
    ("codigo_proveedor",        "Cód. Proveedor",         16, None),
    ("otros_codigos",           "Otros Códigos",          25, None),
]


def generate_excel(document, articles) -> bytes:
    """Generate Excel workbook with article data.

    Args:
        document: Document ORM object
        articles: List of Article ORM objects

    Returns:
        Excel file as bytes
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side, numbers
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl not installed")
        return b""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ARTÍCULOS"

    # Header styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, (field, label, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    # Alternating row styles
    row_fill_even = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    price_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_right = Alignment(horizontal="right", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")

    # Write data rows
    for row_idx, article in enumerate(articles, start=2):
        is_even = (row_idx % 2 == 0)

        # Build row data dict
        otros_codigos_str = ""
        if article.otros_codigos:
            try:
                oc = json.loads(article.otros_codigos) if isinstance(article.otros_codigos, str) else article.otros_codigos
                otros_codigos_str = "; ".join(f"{k}:{v}" for k, v in oc.items()) if oc else ""
            except Exception:
                otros_codigos_str = str(article.otros_codigos)

        row_data = {
            "proveedor": document.supplier_name or "",
            "fecha_doc": document.doc_date.isoformat() if document.doc_date else "",
            "num_doc": document.doc_number or "",
            "descripcion": article.descripcion,
            "cantidad": article.cantidad,
            "precio_unitario_bruto": article.precio_unitario_bruto,
            "descuento_1": (article.descuento_1 / 100) if article.descuento_1 else None,
            "descuento_2": (article.descuento_2 / 100) if article.descuento_2 else None,
            "descuento_3": (article.descuento_3 / 100) if article.descuento_3 else None,
            "descuento_4": (article.descuento_4 / 100) if article.descuento_4 else None,
            "coste_neto_unitario": article.coste_neto_unitario,
            "coste_neto_total": article.coste_neto_total,
            "iva_pct": article.iva_pct / 100,
            "recargo_pct": (article.recargo_pct / 100) if article.recargo_pct else None,
            "margen_pct": article.margen_pct / 100,
            "pvp_sin_iva": article.pvp_sin_iva,
            "pvp_con_iva": article.pvp_con_iva,
            "codigo_principal": article.codigo_principal or "",
            "ean": article.ean or "",
            "codigo_fabricante": article.codigo_fabricante or "",
            "codigo_proveedor": article.codigo_proveedor or "",
            "otros_codigos": otros_codigos_str,
        }

        price_col_names = {
            "precio_unitario_bruto", "coste_neto_unitario", "coste_neto_total",
            "pvp_sin_iva", "pvp_con_iva"
        }

        for col_idx, (field, _, _, num_format) in enumerate(COLUMNS, start=1):
            value = row_data.get(field)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border

            # Background fill
            if field in price_col_names:
                cell.fill = price_fill
            elif is_even:
                cell.fill = row_fill_even

            # Number format
            if num_format and value is not None:
                if "%" in num_format:
                    cell.number_format = "0.00%"
                elif "€" in num_format:
                    cell.number_format = '#,##0.00 "€"'
                else:
                    cell.number_format = num_format

            # Alignment
            if isinstance(value, (int, float)):
                cell.alignment = data_align_right
            elif field in ("descripcion", "otros_codigos"):
                cell.alignment = data_align_left
            else:
                cell.alignment = data_align_center

        ws.row_dimensions[row_idx].height = 18

    # Output to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
