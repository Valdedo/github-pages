"""Tests for Excel and PDF label export services."""
import pytest
from unittest.mock import MagicMock
from datetime import date


def make_mock_document(supplier_name="Test S.L.", doc_number="ALB-001", doc_date=date(2024, 3, 15)):
    doc = MagicMock()
    doc.id = 1
    doc.supplier_name = supplier_name
    doc.doc_number = doc_number
    doc.doc_date = doc_date
    doc.original_filename = "test_albaran.pdf"
    return doc


def make_mock_article(
    id=1, line_number=1, descripcion="Tornillo M6x25",
    cantidad=10.0, precio_unitario_bruto=5.50,
    descuento_1=30.0, descuento_2=10.0, descuento_3=None, descuento_4=None,
    coste_neto_unitario=3.465, coste_neto_total=34.65,
    iva_pct=21.0, recargo_pct=None,
    margen_pct=70.0, pvp_sin_iva=5.89, pvp_con_iva=7.13,
    codigo_principal="FER-001", ean="8414702123456",
    codigo_fabricante="STL-001", codigo_proveedor=None,
    otros_codigos=None, product_info_id=None,
):
    art = MagicMock()
    art.id = id
    art.line_number = line_number
    art.descripcion = descripcion
    art.cantidad = cantidad
    art.precio_unitario_bruto = precio_unitario_bruto
    art.descuento_1 = descuento_1
    art.descuento_2 = descuento_2
    art.descuento_3 = descuento_3
    art.descuento_4 = descuento_4
    art.coste_neto_unitario = coste_neto_unitario
    art.coste_neto_total = coste_neto_total
    art.iva_pct = iva_pct
    art.recargo_pct = recargo_pct
    art.margen_pct = margen_pct
    art.pvp_sin_iva = pvp_sin_iva
    art.pvp_con_iva = pvp_con_iva
    art.codigo_principal = codigo_principal
    art.ean = ean
    art.codigo_fabricante = codigo_fabricante
    art.codigo_proveedor = codigo_proveedor
    art.otros_codigos = otros_codigos
    art.product_info_id = product_info_id
    return art


class TestExcelExport:
    def test_generates_bytes(self):
        from app.services.excel_service import generate_excel
        doc = make_mock_document()
        articles = [make_mock_article(i + 1) for i in range(3)]
        result = generate_excel(doc, articles)
        assert isinstance(result, bytes)
        assert len(result) > 1000

    def test_excel_has_correct_columns(self):
        from app.services.excel_service import generate_excel, COLUMNS
        import openpyxl
        import io

        doc = make_mock_document()
        articles = [make_mock_article()]
        result = generate_excel(doc, articles)

        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        headers = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUMNS))]

        # Check key columns are present
        assert "Descripción" in headers
        assert "PVP con IVA" in headers
        assert "EAN/UPC" in headers
        assert "Coste Neto Unit." in headers

    def test_excel_has_data_rows(self):
        from app.services.excel_service import generate_excel
        import openpyxl
        import io

        doc = make_mock_document()
        articles = [make_mock_article(i + 1, line_number=i + 1, descripcion=f"Artículo {i + 1}") for i in range(5)]
        result = generate_excel(doc, articles)

        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        # Row 1 = headers, rows 2+ = data
        assert ws.max_row == 6  # 1 header + 5 articles

    def test_excel_empty_articles(self):
        from app.services.excel_service import generate_excel
        doc = make_mock_document()
        result = generate_excel(doc, [])
        assert isinstance(result, bytes)
        assert len(result) > 0


class TestLabelExport:
    def test_generates_pdf_bytes(self):
        from app.services.label_service import generate_labels_pdf
        articles = [make_mock_article(i + 1) for i in range(3)]
        result = generate_labels_pdf(articles, base_url="http://localhost:3000")
        assert isinstance(result, bytes)
        assert len(result) > 1000
        # PDF magic bytes
        assert result[:4] == b'%PDF'

    def test_empty_articles_still_generates(self):
        from app.services.label_service import generate_labels_pdf
        result = generate_labels_pdf([], base_url="http://localhost:3000")
        assert isinstance(result, bytes)

    def test_ean_barcode_detection(self):
        from app.services.label_service import generate_barcode_image
        # 13-digit EAN
        ean_bytes = generate_barcode_image("8414702123456", is_ean=True)
        assert ean_bytes is not None
        assert len(ean_bytes) > 100

    def test_code128_barcode(self):
        from app.services.label_service import generate_barcode_image
        code128_bytes = generate_barcode_image("FER-001", is_ean=False)
        assert code128_bytes is not None
        assert len(code128_bytes) > 100

    def test_qr_generation(self):
        from app.services.label_service import generate_qr_image
        qr_bytes = generate_qr_image("http://localhost:3000/producto/1")
        assert qr_bytes is not None
        assert len(qr_bytes) > 100


class TestApiEndpoints:
    def test_health_endpoint(self, client):
        resp = client.get("/health")
        assert resp.status_code == 200
        assert resp.json()["status"] == "ok"

    def test_list_documents_empty(self, client):
        resp = client.get("/api/documents")
        assert resp.status_code == 200
        assert resp.json() == []

    def test_get_settings_default(self, client):
        resp = client.get("/api/settings")
        assert resp.status_code == 200
        data = resp.json()
        assert "margin_tiers" in data
        assert len(data["margin_tiers"]) == 6
        assert data["rounding_mode"] == "standard"

    def test_update_settings(self, client):
        resp = client.put("/api/settings", json={
            "rounding_mode": "psychological",
            "rounding_decimals": 2,
        })
        assert resp.status_code == 200
        assert resp.json()["rounding_mode"] == "psychological"

    def test_create_supplier(self, client):
        resp = client.post("/api/settings/suppliers", json={
            "name": "Test Supplier",
            "detection_keywords": ["test", "supplier"],
            "template_config": {"hints": "columna 3 es descuento"},
        })
        assert resp.status_code == 200
        data = resp.json()
        assert data["name"] == "Test Supplier"
        assert "test" in data["detection_keywords"]
